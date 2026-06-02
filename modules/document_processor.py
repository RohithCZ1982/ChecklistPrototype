"""Document ingestion, text extraction, and entity detection."""

import io
import os
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from config import UPLOAD_DIR
from utils.chunker import chunk_by_page, chunk_text
from utils.helpers import clean_text, file_hash, regex_ner


# ─── Text extraction ──────────────────────────────────────────────────────────

def _rects_overlap(a, b) -> bool:
    """True when rectangles (x0,y0,x1,y1) overlap."""
    return not (a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1])


def _deduplicate_md_table(md: str) -> str:
    """
    PyMuPDF sometimes splits merged cells into phantom duplicate columns, e.g.
    Col1/Test/Col3 all containing the same data.  This function:
    - groups columns that have identical data across all rows
    - within each group keeps the column with a meaningful header (not 'ColN')
    - rebuilds the table with only unique columns
    """
    lines = [l for l in md.strip().split("\n") if l.strip()]
    if len(lines) < 2:
        return md

    def _cells(line: str) -> list:
        return [c.strip() for c in line.strip().strip("|").split("|")]

    headers = _cells(lines[0])
    body_rows: list = []
    for line in lines[1:]:
        if re.match(r"^\s*\|?[-:\s|]+\|?\s*$", line):
            continue
        body_rows.append(_cells(line))

    if not body_rows:
        return md

    n = len(headers)
    # Tuple of data values per column (bounded to header count)
    col_data = [
        tuple(row[ci] if ci < len(row) else "" for row in body_rows)
        for ci in range(n)
    ]

    # Map each unique data-tuple to the preferred column index
    _placeholder = re.compile(r"^Col\d+$", re.I)
    data_to_col: dict = {}
    for ci, data in enumerate(col_data):
        if data not in data_to_col:
            data_to_col[data] = ci
        else:
            prev = data_to_col[data]
            prev_hdr = headers[prev] if prev < n else ""
            curr_hdr = headers[ci] if ci < n else ""
            # Prefer the named header over a placeholder
            if _placeholder.match(prev_hdr) and not _placeholder.match(curr_hdr):
                data_to_col[data] = ci

    keep = sorted(set(data_to_col.values()))
    if len(keep) == n:
        return md  # nothing to remove

    def _row(cells: list) -> str:
        return "| " + " | ".join(cells[ci] if ci < len(cells) else "" for ci in keep) + " |"

    sep = "| " + " | ".join(["---"] * len(keep)) + " |"
    return "\n".join([_row(headers), sep] + [_row(r) for r in body_rows])


def _table_caption(page, tab_bbox: tuple, max_gap: float = 60.0) -> str:
    """
    Return the text of the block sitting immediately above tab_bbox.
    Only considers blocks within max_gap points that horizontally overlap the table.
    This captures titles like "TABLE 302B.02 OTHER ACCEPTANCE CRITERIA" which
    PyMuPDF's find_tables() excludes because they sit outside the cell grid.
    """
    table_top = tab_bbox[1]
    best_gap = max_gap + 1.0
    caption = ""
    for blk in page.get_text("blocks", sort=True):
        if blk[6] != 0:          # skip image blocks
            continue
        bx0, by0, bx1, by1 = blk[:4]
        gap = table_top - by1    # positive = block is above table
        if gap <= 0 or gap >= best_gap:
            continue
        # Block must overlap the table horizontally
        if bx1 < tab_bbox[0] or bx0 > tab_bbox[2]:
            continue
        best_gap = gap
        caption = blk[4].strip().replace("\n", " ")
    return caption


def extract_pdf(file_path: Path) -> Tuple[str, List[Tuple[int, str]], int]:
    """
    Returns (full_text, [(page_num, page_text), ...], page_count).
    Tables are extracted as markdown chunks prefixed with their caption/title so
    that searches for the table name (e.g. "TABLE 302B.02") find the right chunk.
    Text blocks that belong to a detected table are excluded from regular-text
    chunks to avoid duplicate flat-text hits.
    """
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(str(file_path))
        pages = []
        for i, page in enumerate(doc, start=1):
            # ── Extract tables as markdown (with caption) ─────────────────────
            table_bboxes: list = []
            try:
                finder = page.find_tables()
                for tab in finder.tables:
                    md = tab.to_markdown()
                    if not md.strip():
                        continue
                    md = _deduplicate_md_table(md)  # remove phantom duplicate columns
                    tab_bbox = tuple(tab.bbox)
                    caption = _table_caption(page, tab_bbox)
                    chunk = f"{caption}\n{md.strip()}" if caption else md.strip()
                    pages.append((i, chunk))
                    table_bboxes.append(tab_bbox)
            except Exception:
                pass  # find_tables not available or no tables found

            # ── Regular text — skip blocks that belong to a detected table ────
            if table_bboxes:
                parts = []
                for block in page.get_text("blocks", sort=True):
                    if block[6] != 0:  # skip image blocks
                        continue
                    if any(_rects_overlap(block[:4], tb) for tb in table_bboxes):
                        continue
                    block_text = block[4].strip()
                    if block_text:
                        parts.append(block_text)
                text = "\n\n".join(parts)
            else:
                text = page.get_text("text")

            if text.strip():
                pages.append((i, clean_text(text)))

        full_text = "\n\n".join(t for _, t in pages)
        return full_text, pages, len(doc)
    except ImportError:
        raise RuntimeError("PyMuPDF not installed. Run: pip install PyMuPDF")
    except Exception as e:
        raise RuntimeError(f"PDF extraction failed: {e}")


def extract_docx(file_path: Path) -> Tuple[str, List[Tuple[int, str]], int]:
    """Extract text from DOCX. Pages are approximated by paragraph groups."""
    try:
        from docx import Document
        doc = Document(str(file_path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Group ~30 paragraphs per virtual page
        page_size = 30
        pages = []
        for i in range(0, max(len(paragraphs), 1), page_size):
            group = paragraphs[i : i + page_size]
            page_text = "\n".join(group)
            pages.append((i // page_size + 1, clean_text(page_text)))

        # Also extract tables as markdown
        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                rows.append(cells)
            if len(rows) >= 2:
                # Build markdown table
                max_cols = max(len(r) for r in rows)
                rows = [r + [""] * (max_cols - len(r)) for r in rows]
                header = "| " + " | ".join(rows[0]) + " |"
                separator = "| " + " | ".join(["---"] * max_cols) + " |"
                body = "\n".join("| " + " | ".join(r) + " |" for r in rows[1:])
                md_table = f"{header}\n{separator}\n{body}"
                pages.append((len(pages) + 1, md_table))

        full_text = clean_text("\n\n".join(t for _, t in pages))
        return full_text, pages, len(pages)
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")
    except Exception as e:
        raise RuntimeError(f"DOCX extraction failed: {e}")


def extract_image(file_path: Path) -> Tuple[str, List[Tuple[int, str]], int]:
    """OCR image using pytesseract."""
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(str(file_path))
        text = pytesseract.image_to_string(img)
        text = clean_text(text)
        return text, [(1, text)], 1
    except ImportError:
        raise RuntimeError("pytesseract or Pillow not installed.")
    except Exception as e:
        raise RuntimeError(f"Image OCR failed: {e}")


def extract_txt(file_path: Path) -> Tuple[str, List[Tuple[int, str]], int]:
    for enc in ("utf-8", "latin-1", "cp1252"):
        try:
            text = file_path.read_text(encoding=enc)
            text = clean_text(text)
            # Split into virtual pages of ~3000 chars
            pages = []
            chunk_sz = 3000
            for i in range(0, max(len(text), 1), chunk_sz):
                pages.append((i // chunk_sz + 1, text[i : i + chunk_sz]))
            return text, pages, len(pages)
        except UnicodeDecodeError:
            continue
    raise RuntimeError("Could not decode text file.")


def extract_text(file_path: Path, file_type: str) -> Tuple[str, List[Tuple[int, str]], int]:
    """Dispatch extraction based on file type."""
    ft = file_type.lower().lstrip(".")
    if ft == "pdf":
        return extract_pdf(file_path)
    if ft == "docx":
        return extract_docx(file_path)
    if ft in ("png", "jpg", "jpeg", "tiff", "bmp", "gif"):
        return extract_image(file_path)
    if ft == "txt":
        return extract_txt(file_path)
    raise ValueError(f"Unsupported file type: {ft}")


# ─── Metadata inference ───────────────────────────────────────────────────────

_DOC_TYPE_PATTERNS = {
    "Invoice": re.compile(r'\binvoice\b', re.I),
    "Contract": re.compile(r'\b(contract|agreement|terms)\b', re.I),
    "Report": re.compile(r'\b(report|analysis|summary|review)\b', re.I),
    "Resume": re.compile(r'\b(resume|curriculum vitae|cv|work experience)\b', re.I),
    "Legal": re.compile(r'\b(plaintiff|defendant|court|law|clause|attorney)\b', re.I),
    "Financial": re.compile(r'\b(balance sheet|profit|loss|revenue|financial statement)\b', re.I),
    "Medical": re.compile(r'\b(patient|diagnosis|prescription|medical|clinical)\b', re.I),
    "Research": re.compile(r'\b(abstract|methodology|hypothesis|findings|literature)\b', re.I),
}


def infer_doc_type(text: str) -> str:
    for label, pattern in _DOC_TYPE_PATTERNS.items():
        if pattern.search(text[:3000]):
            return label
    return "General"


def infer_title(text: str, filename: str) -> str:
    """Use first non-empty line or filename stem as title."""
    for line in text.split("\n"):
        line = line.strip()
        if len(line) > 5 and len(line) < 200:
            return line
    return Path(filename).stem.replace("_", " ").replace("-", " ").title()


# ─── Main processor ───────────────────────────────────────────────────────────

def process_document(
    file_bytes: bytes,
    filename: str,
    doc_id: str,
) -> Dict:
    """
    Full pipeline: save → extract → chunk → entities.

    Returns a dict with keys:
      text, pages, chunks, entities, metadata
    """
    file_type = Path(filename).suffix.lstrip(".").lower()
    save_path = UPLOAD_DIR / f"{doc_id}_{filename}"

    # Persist the file
    save_path.write_bytes(file_bytes)

    # Extract text
    full_text, pages, page_count = extract_text(save_path, file_type)

    # Build chunks (preserve page provenance)
    chunks = chunk_by_page(pages)

    # Lightweight NER on first 8 000 chars
    entities = regex_ner(full_text[:8000])

    # Attempt spaCy NER if available
    spacy_entities = _spacy_ner(full_text[:8000])
    entities.extend(spacy_entities)

    # Deduplicate entities by (type, text)
    seen = set()
    deduped_entities = []
    for e in entities:
        key = (e["type"], e["text"].lower())
        if key not in seen:
            seen.add(key)
            deduped_entities.append(e)

    metadata = {
        "title": infer_title(full_text, filename),
        "doc_type": infer_doc_type(full_text),
        "page_count": page_count,
        "char_count": len(full_text),
        "word_count": len(full_text.split()),
        "chunk_count": len(chunks),
    }

    return {
        "text": full_text,
        "pages": pages,
        "chunks": chunks,
        "entities": deduped_entities,
        "metadata": metadata,
        "save_path": str(save_path),
    }


def _spacy_ner(text: str) -> List[Dict]:
    """Optional spaCy NER – silently skipped if spaCy/model not available."""
    try:
        import spacy
        try:
            nlp = spacy.load("en_core_web_sm")
        except OSError:
            return []
        doc = nlp(text[:5000])
        results = []
        seen = set()
        for ent in doc.ents:
            key = (ent.label_, ent.text.lower())
            if key not in seen:
                seen.add(key)
                results.append({"type": ent.label_, "text": ent.text, "confidence": 0.80})
        return results
    except ImportError:
        return []


def get_document_file_path(doc_id: str, filename: str) -> Optional[Path]:
    """Reconstruct the saved file path for a document."""
    path = UPLOAD_DIR / f"{doc_id}_{filename}"
    return path if path.exists() else None


def detect_similarity(text1: str, text2: str) -> float:
    """
    Quick Jaccard similarity between two texts based on word sets.
    Returns value in [0, 1].
    """
    words1 = set(re.findall(r'\b\w{3,}\b', text1.lower()))
    words2 = set(re.findall(r'\b\w{3,}\b', text2.lower()))
    if not words1 or not words2:
        return 0.0
    return len(words1 & words2) / len(words1 | words2)
