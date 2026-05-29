"""Export template results to JSON, Excel, and Word formats."""

from __future__ import annotations

import io
import json
from datetime import datetime
from typing import Any, Dict, List, Optional


# ─── JSON ─────────────────────────────────────────────────────────────────────

def export_json(data: Dict) -> bytes:
    """Serialize result dict to formatted JSON bytes."""
    return json.dumps(data, indent=2, default=str, ensure_ascii=False).encode("utf-8")


# ─── Excel ────────────────────────────────────────────────────────────────────

def export_excel(template_output: Dict) -> bytes:
    """
    Build a multi-sheet Excel workbook from template output.
    Sheet 1: Summary (field → synthesized answer + confidence)
    Sheet 2+: Per-field hit details
    """
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")

    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2A5298")
    wrap = Alignment(wrap_text=True, vertical="top")

    # ── Summary sheet ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    template_name = template_output.get("template", {}).get("name", "Template")
    ws["A1"] = f"Template: {template_name}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    ws.append([])

    headers = ["Field", "Synthesized Answer", "Top Confidence", "Hits Found"]
    ws.append(headers)
    for col, _ in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    fields = template_output.get("fields", {})
    for row_i, (fname, fdata) in enumerate(fields.items(), start=5):
        ws.cell(row=row_i, column=1, value=fname)
        ws.cell(row=row_i, column=2, value=fdata.get("synthesized_answer", "")).alignment = wrap
        ws.cell(row=row_i, column=3, value=round(fdata.get("top_confidence", 0), 3))
        ws.cell(row=row_i, column=4, value=len(fdata.get("hits", [])))

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 12

    # ── Per-field detail sheets ───────────────────────────────────────────────
    for fname, fdata in fields.items():
        sheet_name = fname[:31].replace("/", "-").replace("\\", "-")
        ws2 = wb.create_sheet(title=sheet_name)
        ws2.append(["#", "Document", "Page", "Confidence", "Source", "Text Snippet"])
        for col in range(1, 7):
            cell = ws2.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill

        for i, hit in enumerate(fdata.get("hits", []), start=1):
            ws2.append([
                i,
                hit.get("document_id", "")[:40],
                hit.get("page_number", ""),
                round(hit.get("confidence", 0), 3),
                hit.get("source", ""),
                hit.get("text", "")[:500],
            ])
            ws2.cell(row=i + 1, column=6).alignment = wrap

        ws2.column_dimensions["B"].width = 40
        ws2.column_dimensions["F"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── Word ─────────────────────────────────────────────────────────────────────

def export_word(template_output: Dict) -> bytes:
    """Generate a Word document from template output."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")

    doc = Document()

    # Title
    template_name = template_output.get("template", {}).get("name", "Template Output")
    title = doc.add_heading(template_name, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
    ).runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    if template_output.get("template", {}).get("description"):
        doc.add_paragraph(template_output["template"]["description"])

    doc.add_heading("Extracted Fields", level=1)

    fields = template_output.get("fields", {})
    for fname, fdata in fields.items():
        # Field heading
        h = doc.add_heading(fname, level=2)

        # Confidence badge
        conf = fdata.get("top_confidence", 0)
        conf_text = f"Confidence: {conf:.0%}"
        p = doc.add_paragraph()
        run = p.add_run(conf_text)
        run.bold = True
        run.font.color.rgb = (
            RGBColor(0x00, 0x80, 0x00) if conf >= 0.75
            else RGBColor(0xFF, 0x80, 0x00) if conf >= 0.5
            else RGBColor(0xCC, 0x00, 0x00)
        )

        # Synthesized answer
        answer = fdata.get("synthesized_answer", "")
        if answer:
            doc.add_paragraph(answer)

        # Source snippets (top 3)
        hits = fdata.get("hits", [])[:3]
        if hits:
            doc.add_heading("Supporting Evidence", level=3)
            for i, hit in enumerate(hits, 1):
                snippet = hit.get("text", "")[:400]
                p2 = doc.add_paragraph(style="List Number")
                p2.add_run(f"[Page {hit.get('page_number', '?')}] ").bold = True
                p2.add_run(snippet)

        doc.add_paragraph()  # spacing

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─── Helper: flatten for CSV ─────────────────────────────────────────────────

def export_csv(template_output: Dict) -> bytes:
    """Simple CSV: field, answer, confidence."""
    import csv

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["Field", "Synthesized Answer", "Confidence", "Hits"])
    for fname, fdata in template_output.get("fields", {}).items():
        writer.writerow([
            fname,
            fdata.get("synthesized_answer", ""),
            round(fdata.get("top_confidence", 0), 3),
            len(fdata.get("hits", [])),
        ])
    return buf.getvalue().encode("utf-8")
